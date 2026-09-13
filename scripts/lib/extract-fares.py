"""Read official fare files into JSON. Never evaluates spreadsheet formulas.

Invoked by sync-fares.ts; Python dependencies: openpyxl, xlrd, pdfplumber.
"""
import csv
import io
import json
import re
import sys
import zipfile
from html.parser import HTMLParser
from pathlib import Path


def workbook(data):
    if data[:2] == b"PK":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [(sheet.title, list(sheet.values)) for sheet in book]
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    return [(sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]) for sheet in book.sheets()]


def compact(value):
    return re.sub(r"\s+", "", str(value or "")).replace("EXPO", "엑스포")


def korail(data, family):
    tables = []
    for title, rows in workbook(data):
        fares = []
        if family == "ktx":
            # The first standard column is the CURRENT fare (A), not the future B.
            headers = next((r for r in rows[:10] if compact(r[3]).startswith("일반실")), None)
            if headers is None:
                raise ValueError(f"Unrecognized KTX fare header: {title}")
            first = next((i + 2 for i, cell in enumerate(headers) if compact(cell) == "특실"), None)
            for row in rows:
                if len(row) < 4 or not isinstance(row[3], (int, float)):
                    continue
                if not all(isinstance(row[i], str) and re.search("[가-힣]", row[i]) for i in [1, 2]):
                    continue
                fares.append([compact(row[1]), compact(row[2]), row[3],
                              row[first] if first is not None and isinstance(row[first], (int, float)) else None])
        else:
            # Some sheets have two tables side by side. Read only columns explicitly
            # headed origin/destination/fare; a numeric note is never a fare row.
            columns = set()
            for row in rows:
                for i in range(len(row) - 2):
                    if [compact(x) for x in row[i:i+3]] == ["출발역", "도착역", "운임"]:
                        columns.add(i)
            if not columns:
                raise ValueError(f"Unrecognized regular fare header: {title}")
            for row in rows:
                for i in sorted(columns):
                    if i + 2 < len(row) and isinstance(row[i+2], (int, float)) and all(
                        isinstance(row[j], str) and re.search("[가-힣]", row[j]) for j in [i, i+1]
                    ):
                        fares.append([compact(row[i]), compact(row[i+1]), row[i+2], None])
        if not fares or any(r[2] <= 0 for r in fares):
            raise ValueError(f"Missing/invalid fares: {title}")
        tables.append({"title": title, "family": family, "rows": fares})
    return tables


class Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables, self.table, self.row, self.cell = [], None, None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table": self.table = []
        elif tag == "tr" and self.table is not None: self.row = []
        elif tag in ["th", "td"] and self.row is not None: self.cell = []

    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ["td", "th"] and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def extract(kind, data):
    if kind == "json": return json.loads(data)
    if kind in ["ktx", "saemaeul", "mugunghwa"]: return korail(data, kind)
    if kind == "csv": return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    if kind == "mbta":
        archive = zipfile.ZipFile(io.BytesIO(data))
        # Only fare rules and station/route identities, never trips or stop_times.
        names = ["fare_products", "fare_leg_rules", "fare_transfer_rules", "fare_media", "areas",
                 "stop_areas", "routes", "stops", "feed_info", "timeframes", "calendar", "calendar_dates"]
        result = {name: extract("csv", archive.read(name + ".txt")) for name in names}
        result["calendar"] = [r for r in result["calendar"] if r["service_id"] == "fare_regular"]
        result["calendar_dates"] = [r for r in result["calendar_dates"] if r["service_id"] == "fare_regular"]
        result["stops"] = [{k: r[k] for k in ["stop_id", "stop_name", "parent_station"]} for r in result["stops"]]
        result["routes"] = [{k: r[k] for k in ["route_id", "route_short_name", "route_long_name", "route_type", "network_id"]} for r in result["routes"]]
        rail = {r["route_id"] for r in result["routes"] if r["route_type"] in ["0", "1", "2"]}
        trips = {r["trip_id"]: r["route_id"] for r in extract("csv", archive.read("trips.txt")) if r["route_id"] in rail}
        stops = {r["stop_id"]: r["parent_station"] or r["stop_id"] for r in result["stops"]}
        pairs = set()
        with archive.open("stop_times.txt") as stream:
            for r in csv.DictReader(io.TextIOWrapper(stream, encoding="utf-8-sig")):
                if r["trip_id"] in trips:
                    pairs.add((trips[r["trip_id"]], stops[r["stop_id"]]))
        result["route_stations"] = [{"route_id": route, "station_id": station} for route, station in sorted(pairs)]
        return result
    if kind == "html-tables":
        parser = Tables()
        parser.feed(data.decode("utf-8-sig"))
        return parser.tables
    if kind == "pdf-tables":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            # Retain factual numeric rows, not entire copyrighted conditions.
            result = []
            for i, page in enumerate(pdf.pages):
                for table in page.extract_tables():
                    rows = [row for row in table if all(len(c or "") < 100 for c in row)
                            and sum(bool(re.fullmatch(r"[\d., –-]+", c or "")) for c in row) >= 2]
                    if rows: result.append({"page": i+1, "rows": rows})
            return result
    raise ValueError(f"Unknown fare format: {kind}")


if __name__ == "__main__":
    value = extract(sys.argv[1], Path(sys.argv[2]).read_bytes())
    Path(sys.argv[3]).write_text(json.dumps(value, ensure_ascii=False, separators=(",", ":")))
